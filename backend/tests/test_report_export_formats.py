"""Tests for report export formats: XLSX and PDF (v1.0.1 hotfix).

Validates that Candidate Report and Wetlab Validation Report endpoints
support json, markdown/md, csv, xlsx, and pdf formats.
Also verifies scientific boundary statements are present in all exports.
"""

from __future__ import annotations

import io

import pytest
import pytest_asyncio
from fastapi import status
from httpx import ASGITransport, AsyncClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.crud.projects import create_project
from app.crud.stamp_candidates import create_stamp_candidate, create_stamp_generation_run
from app.database import Base, get_db
from app.main import create_app
from app.schemas.project import ProjectCreate
from app.schemas.stamp import StampCandidateCreate, StampGenerationRunCreate

TEST_DB_URL = "sqlite:///:memory:"
_test_engine = create_engine(
    TEST_DB_URL, echo=False, future=True, connect_args={"check_same_thread": False}, poolclass=StaticPool
)


@pytest.fixture(scope="function")
def db_session():
    Base.metadata.create_all(bind=_test_engine)
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=_test_engine)
    session = SessionLocal()
    yield session
    session.close()
    Base.metadata.drop_all(bind=_test_engine)


@pytest.fixture
def test_project(db_session):
    return create_project(db_session, ProjectCreate(name="Format Export Test Project"))


@pytest.fixture
def test_generation_run(db_session, test_project):
    return create_stamp_generation_run(
        db_session, StampGenerationRunCreate(project_id=test_project.id)
    )


@pytest.fixture
def test_candidate(db_session, test_project, test_generation_run):
    return create_stamp_candidate(
        db_session,
        StampCandidateCreate(
            project_id=test_project.id,
            generation_run_id=test_generation_run.id,
            targeting_peptide_seq="MKWVTFISLL",
            linker_seq="GGGGS",
            full_sequence="MKWVTFISLLGGGGS",
            composite_score=0.85,
            metrics={
                "source": "pepmlm_650m_real",
                "structure_prediction": {
                    "mean_plddt": 65.0,
                    "prediction_status": "COMPUTATIONAL_STRUCTURE_PREDICTION_ONLY",
                },
                "interface_quality": {
                    "pdockq": 0.55,
                    "prediction_status": "COMPUTATIONAL_INTERFACE_QUALITY_PREDICTION_ONLY",
                },
                "energy_quality": {
                    "interaction_energy_kcal_mol": 12.3,
                    "prediction_status": "COMPUTATIONAL_INTERACTION_ENERGY_ESTIMATE_ONLY",
                },
            },
        ),
    )


@pytest.fixture
def test_app(db_session):
    Base.metadata.create_all(bind=db_session.bind)
    app = create_app()

    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    return app


@pytest_asyncio.fixture
async def async_client(test_app):
    transport = ASGITransport(app=test_app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        yield client


# ---------------------------------------------------------------------------
# Candidate Report — API format tests
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_candidate_report_json(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/projects/{test_project.id}/candidate-report?format=json&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    data = resp.json()
    assert data["code"] == 200
    assert data["data"]["report_type"] == "STAMP_CANDIDATE_REPORT"


@pytest.mark.asyncio
async def test_candidate_report_markdown(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/projects/{test_project.id}/candidate-report?format=markdown&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    assert "text/markdown" in resp.headers.get("content-type", "")
    body = resp.text
    assert "STAMP Candidate Report" in body
    assert "NOT_EXPERIMENTALLY_VALIDATED" in body


@pytest.mark.asyncio
async def test_candidate_report_md_alias(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/projects/{test_project.id}/candidate-report?format=md&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    assert "text/markdown" in resp.headers.get("content-type", "")


@pytest.mark.asyncio
async def test_candidate_report_csv(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/projects/{test_project.id}/candidate-report?format=csv&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    assert "text/csv" in resp.headers.get("content-type", "")
    body = resp.text
    assert "rank,candidate_id" in body.lower() or "candidate_id" in body


@pytest.mark.asyncio
async def test_candidate_report_xlsx(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/projects/{test_project.id}/candidate-report?format=xlsx&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    ct = resp.headers.get("content-type", "")
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in ct
    assert len(resp.content) > 0
    # Verify it is a valid ZIP (XLSX) by checking PK header
    assert resp.content[:2] == b"PK"
    # Verify openpyxl can load it
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Summary" in wb.sheetnames
    assert "Candidates" in wb.sheetnames
    assert "Scientific Boundary" in wb.sheetnames


@pytest.mark.asyncio
async def test_candidate_report_pdf(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/projects/{test_project.id}/candidate-report?format=pdf&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    ct = resp.headers.get("content-type", "")
    assert "application/pdf" in ct
    assert len(resp.content) > 0
    assert resp.content[:4] == b"%PDF"


@pytest.mark.asyncio
async def test_candidate_report_invalid_format(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/projects/{test_project.id}/candidate-report?format=xml&top_k=10"
    )
    assert resp.status_code == status.HTTP_422_UNPROCESSABLE_ENTITY


# ---------------------------------------------------------------------------
# Wetlab Validation Report — API format tests
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_wetlab_report_json(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/experimental-validation/projects/{test_project.id}/wetlab-validation-report?format=json&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    data = resp.json()
    assert data["code"] == 200
    assert data["data"]["report_type"] == "STAMP_WETLAB_VALIDATION_REPORT"


@pytest.mark.asyncio
async def test_wetlab_report_markdown(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/experimental-validation/projects/{test_project.id}/wetlab-validation-report?format=markdown&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    assert "text/markdown" in resp.headers.get("content-type", "")
    body = resp.text
    assert "STAMP Wet-lab Validation Report" in body
    # Scientific boundary statement is present even when no candidates have data
    assert "SHORTLISTED" in body or "NOT_EXPERIMENTALLY_VALIDATED" in body or "experimental validation" in body.lower()


@pytest.mark.asyncio
async def test_wetlab_report_md_alias(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/experimental-validation/projects/{test_project.id}/wetlab-validation-report?format=md&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    assert "text/markdown" in resp.headers.get("content-type", "")


@pytest.mark.asyncio
async def test_wetlab_report_csv(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/experimental-validation/projects/{test_project.id}/wetlab-validation-report?format=csv&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    assert "text/csv" in resp.headers.get("content-type", "")
    body = resp.text
    assert "candidate_id" in body.lower()


@pytest.mark.asyncio
async def test_wetlab_report_xlsx(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/experimental-validation/projects/{test_project.id}/wetlab-validation-report?format=xlsx&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    ct = resp.headers.get("content-type", "")
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in ct
    assert len(resp.content) > 0
    assert resp.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Summary" in wb.sheetnames
    assert "Candidates" in wb.sheetnames
    assert "Scientific Boundary" in wb.sheetnames


@pytest.mark.asyncio
async def test_wetlab_report_pdf(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/experimental-validation/projects/{test_project.id}/wetlab-validation-report?format=pdf&top_k=10"
    )
    assert resp.status_code == status.HTTP_200_OK
    ct = resp.headers.get("content-type", "")
    assert "application/pdf" in ct
    assert len(resp.content) > 0
    assert resp.content[:4] == b"%PDF"


@pytest.mark.asyncio
async def test_wetlab_report_invalid_format(async_client, test_project):
    resp = await async_client.get(
        f"/api/v1/experimental-validation/projects/{test_project.id}/wetlab-validation-report?format=xml&top_k=10"
    )
    assert resp.status_code == status.HTTP_400_BAD_REQUEST


# ---------------------------------------------------------------------------
# Scientific boundary tests for binary exports
# ---------------------------------------------------------------------------

def test_candidate_report_xlsx_contains_boundary(db_session, test_project, test_candidate):
    from app.services.candidate_report_export import build_candidate_report_data, render_candidate_report_xlsx
    report_data = build_candidate_report_data(db_session, test_project.id, top_k=10)
    xlsx_bytes = render_candidate_report_xlsx(report_data)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Scientific Boundary"]
    cell_text = " ".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value)
    assert "NOT experimentally validated" in cell_text or "computational" in cell_text.lower()


def test_candidate_report_pdf_contains_boundary(db_session, test_project, test_candidate):
    from app.services.candidate_report_export import build_candidate_report_data, render_candidate_report_pdf
    report_data = build_candidate_report_data(db_session, test_project.id, top_k=10)
    pdf_bytes = render_candidate_report_pdf(report_data)
    # Extract text using pdfminer
    from pdfminer.high_level import extract_text
    text = extract_text(io.BytesIO(pdf_bytes))
    assert "Scientific Boundary" in text
    # PDF text may have line breaks inside phrases; normalize whitespace
    normalized = " ".join(text.split())
    assert "NOT experimentally validated" in normalized


def test_wetlab_report_xlsx_contains_boundary(db_session, test_project, test_candidate):
    from app.services.wetlab_report_export import build_wetlab_validation_report_data, render_wetlab_report_xlsx
    report_data = build_wetlab_validation_report_data(db_session, test_project.id, top_k=10)
    xlsx_bytes = render_wetlab_report_xlsx(report_data)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Scientific Boundary"]
    cell_text = " ".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value)
    assert "SHORTLISTED" in cell_text or "NOT_EXPERIMENTALLY_VALIDATED" in cell_text or "experimental" in cell_text.lower()


def test_wetlab_report_pdf_contains_boundary(db_session, test_project, test_candidate):
    from app.services.wetlab_report_export import build_wetlab_validation_report_data, render_wetlab_report_pdf
    report_data = build_wetlab_validation_report_data(db_session, test_project.id, top_k=10)
    pdf_bytes = render_wetlab_report_pdf(report_data)
    from pdfminer.high_level import extract_text
    text = extract_text(io.BytesIO(pdf_bytes))
    assert "Scientific Boundary" in text
    assert "SHORTLISTED" in text or "NOT_EXPERIMENTALLY_VALIDATED" in text


# ---------------------------------------------------------------------------
# Fabrication guard tests (regression)
# ---------------------------------------------------------------------------

def test_candidate_report_no_fabricated_mic_mbc(db_session, test_project, test_candidate):
    from app.services.candidate_report_export import build_candidate_report_data, validate_report_no_fabricated_metrics
    report_data = build_candidate_report_data(db_session, test_project.id, top_k=10)
    validate_report_no_fabricated_metrics(report_data)
    for c in report_data["candidates"]:
        assert "MIC" not in c
        assert "MBC" not in c
        assert "hemolysis" not in c
        assert "toxicity" not in c
        assert c["validation_status"] == "NOT_EXPERIMENTALLY_VALIDATED"


def test_candidate_report_no_official_mmgbsa_delta_g(db_session, test_project, test_candidate):
    from app.services.candidate_report_export import build_candidate_report_data
    report_data = build_candidate_report_data(db_session, test_project.id, top_k=10)
    for c in report_data["candidates"]:
        eq = c.get("energy_quality", {})
        assert eq.get("not_mmgbsa_delta_g") is True
        assert eq.get("not_docking_score") is True


def test_wetlab_report_no_fabricated_measurements(db_session, test_project, test_candidate):
    from app.services.wetlab_report_export import build_wetlab_validation_report_data
    report_data = build_wetlab_validation_report_data(db_session, test_project.id, top_k=10)
    for c in report_data["candidates"]:
        assert c["measurement_count"] == 0
        assert c["experimental_measurements"] == []
        assert c["validation_status"] == "NOT_EXPERIMENTALLY_VALIDATED"
